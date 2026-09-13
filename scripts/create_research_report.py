from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "research"
OUT.mkdir(exist_ok=True)
OUTPUT = OUT / "image_to_brick_model_research.docx"


SOURCES = {
    1: ("OpenAI", "GPT 6 Astra Model", "2026", "https://developers.openai.com/api/docs/models/gpt-6-astra"),
    2: ("OpenAI", "Architectural visualization with Astra", "4 September 2026", "https://developers.openai.com/blog/architectural-visualization-with-astra"),
    3: ("LDraw.org", "LDraw Parts Library", "accessed September 2026", "https://library.ldraw.org/"),
    4: ("LDraw.org Standards Board", "LDraw File Format Specification", "29 April 2012", "https://www.ldraw.org/article/218.html"),
    5: ("LDraw.org", "Contributor Agreement", "accessed September 2026", "https://ldraw.org/ldraw-org-contributor-agreement"),
    6: ("Roland Melkert", "LDCad Shadow Library", "accessed September 2026", "https://www.melkert.net/LDCad/tech/shadowLib"),
    7: ("BrickLink Studio Help Center", "Sculpture", "accessed September 2026", "https://studiohelp.bricklink.com/hc/en-us/articles/6508264220183-Sculpture"),
    8: ("BrickLink Studio Help Center", "Exporting to other formats", "accessed September 2026", "https://studiohelp.bricklink.com/hc/en-us/articles/6502197862679-Exporting-to-other-formats"),
    9: ("Rebrickable", "API Documentation", "accessed September 2026", "https://rebrickable.com/api/v3/docs/"),
    10: ("BrickLink", "API Manual", "accessed September 2026", "https://static.bricklink.com/alpha/default/api_wiki.html"),
    11: ("R Testuz, Y Schwartzburg, and M Pauly", "Automatic Generation of Constructable Brick Sculptures", "2013", "https://diglib.eg.org/items/82e2ee53-fa57-4929-ac81-85acd0d9f855"),
    12: ("S Luo et al", "Legolization Optimizing LEGO Designs", "2015", "https://www.cs.columbia.edu/~yonghao/siga15/luo-Legolization.pdf"),
    13: ("K Lennon et al", "Image2Lego Customized LEGO Set Generation from Images", "2021", "https://arxiv.org/abs/2108.08477"),
    14: ("A Pun et al", "Generating Physically Stable and Buildable LEGO Designs from Text", "2025", "https://arxiv.org/abs/2505.05469"),
    15: ("Z Ni et al", "BrickAnything Geometry Conditioned Buildable Brick Generation", "2026", "https://arxiv.org/abs/2605.26182"),
    16: ("Meta AI", "SAM 3D Objects", "2025", "https://github.com/facebookresearch/sam-3d-objects"),
    17: ("Microsoft", "TRELLIS 2", "2026", "https://github.com/microsoft/TRELLIS.2"),
    18: ("COLMAP", "Structure from Motion and Multi View Stereo", "accessed September 2026", "https://colmap.github.io/"),
    19: ("The LEGO Group", "Fair Play", "accessed September 2026", "https://www.lego.com/en-us/legal/notices-and-policies/fair-play"),
    20: ("J Johnson", "BrickBuilder AI", "2026", "https://github.com/jjohnson5253/brickbuilderai"),
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)
        borders.append(tag)


def add_hyperlink(paragraph, text, url, color="1F4E79", underline=True):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_citations(paragraph, refs):
    if not refs:
        return
    run = paragraph.add_run(" [" + ", ".join(str(r) for r in refs) + "]")
    run.font.superscript = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(70, 70, 70)


def add_para(doc, text="", refs=None, bold_lead=None, style=None, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    add_citations(p, refs)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc, text, refs=None, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = add_para(doc, text, refs=refs, style=style)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_number(doc, number, text, refs=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.34)
    p.paragraph_format.first_line_indent = Inches(-0.24)
    p.add_run(f"{number}. ").bold = True
    p.add_run(text)
    add_citations(p, refs)
    p.paragraph_format.space_after = Pt(4)
    return p


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    old = p_pr.find(qn("w:pBdr"))
    if old is not None:
        p_pr.remove(old)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    p_pr.append(borders)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        if widths:
            cell.width = Inches(widths[idx])
        set_cell_shading(cell, "23395B")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            if widths:
                cell.width = Inches(widths[idx])
            if r_idx % 2:
                set_cell_shading(cell, "F3F6FA")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
            if idx > 0 and len(str(value)) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(31, 31, 31)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(28)
styles["Title"].font.bold = True
styles["Heading 1"].font.size = Pt(17)
styles["Heading 1"].font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(7)
styles["Heading 2"].font.size = Pt(13)
styles["Heading 2"].font.bold = True
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(5)
styles["Heading 3"].font.size = Pt(11)
styles["Heading 3"].font.bold = True

header = section.header.paragraphs[0]
header.text = "Image to buildable brick model research"
header.style = styles["Normal"]
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(95, 95, 95)
footer = section.footer.paragraphs[0]
add_page_number(footer)
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(95, 95, 95)

title = doc.add_paragraph(style="Title")
title.add_run("Image to Buildable Brick Model Architecture")
remove_paragraph_borders(title)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
r = subtitle.add_run("Research assessment and recommended MVP pipeline")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(55, 55, 55)

add_para(doc, "This report evaluates four ways to turn photographs of an everyday object into a physically buildable brick sculpture and a valid parts list. It covers available digital part libraries, current image-to-3D systems, computational brick-layout methods, OpenAI GPT 6 Astra’s role, and a staged implementation plan.")
add_para(doc, "Recommendation. Build a hybrid pipeline: reconstruct a coarse, uncertainty-aware 3D target from guided images; voxelize it on a stud-and-plate lattice; then solve for an assembly made only from a curated catalog of real parts. Use Astra as the multimodal planner, tool-using orchestrator, and visual critic. Do not treat Astra as a native 3D reconstruction endpoint or allow it to emit an unchecked final assembly.", bold_lead="Recommendation.")

add_heading(doc, "Executive decision", 1)
add_para(doc, "The strongest route is an Approach 4 that combines the useful parts of Approaches 1 and 3. A 3D intermediate is valuable, but it should be a design target—occupancy, silhouettes, landmarks, surface colors, and uncertainty—not necessarily a polished production mesh. The final model should be created by a constrained assembly engine operating over real part definitions.")
add_para(doc, "This separation matters because image reconstruction and brick design optimize different things. A plausible mesh can have invented hidden geometry. A visually accurate brick shell can collapse. A collection of connected bricks can be stable but fail to resemble the object. The system needs independent representations and tests for appearance, part legality, connectivity, stability, availability, and assembly order.")

add_table(doc,
    ["Option", "Core idea", "Strength", "Main failure", "Verdict"],
    [
        ["1", "Photos to 3D mesh to brick layout", "Clear stages and strong baselines", "A detailed mesh creates false precision and does not solve buildability", "Use a coarse version"],
        ["2", "Generate a brick assembly directly", "Potentially expressive and fast", "Unreliable geometry, legality, stability, and inventory without hard validation", "Research track only"],
        ["3", "Search the full part library", "Uses authentic parts from the start", "Combinatorial explosion and incomplete connector metadata", "Use a curated subset"],
        ["4", "Uncertain 3D target plus catalog constrained solver", "Balances fidelity, legality, cost, and stability", "Requires a purpose-built optimizer and evaluation harness", "Recommended"],
    ], widths=[0.48, 1.48, 1.55, 2.2, 0.95], font_size=8.5)

add_heading(doc, "Why the problem is harder than image to 3D", 1)
add_para(doc, "The output is not just a render. It is a discrete, physical program: every element needs a part ID, color, pose, legal connection, non-colliding volume, support path, and useful position in an assembly sequence. A user also expects the listed part-color combinations to exist and to be obtainable at a tolerable price.")
add_bullet(doc, "Geometry is ambiguous. A single image cannot determine the back, underside, true scale, or hidden cavities. Even multiple images can fail on reflective, transparent, furry, thin, or moving subjects.")
add_bullet(doc, "Brick geometry is anisotropic. The natural grid is one stud by one stud horizontally and one plate vertically; a normal brick spans three plate layers. Treating cells as cubic distorts proportions.")
add_bullet(doc, "Buildability is global. Greedily combining 1 by 1 cells reduces part count but can align seams between layers, create articulation points, strand islands, or leave overhangs that depend on excessive clutch force.", refs=[11, 12])
add_bullet(doc, "Visual identity is nonuniform. A cat’s ears, eye spacing, muzzle, tail gesture, and coat patches matter more than equal error everywhere. A bottle’s silhouette, neck, cap, and label band dominate recognition.")
add_bullet(doc, "Availability is a part-color-time constraint. A geometrically valid part can be unavailable or very expensive in a desired color; substitutions can materially change the design.", refs=[9, 10])

add_heading(doc, "What is available now", 1)
add_heading(doc, "Digital parts and interchange formats", 2)
add_para(doc, "LDraw is the best foundation for the canonical geometry and output format. Its current official library reports 17,116 unique shapes or patterned parts, supplies a text-based model format with positioned part references and transformations, and continues to receive frequent updates. The library’s contributor agreement supports CC BY 4.0 or CC0 for submitted work, but attribution and per-file license handling still need to be implemented carefully.", refs=[3, 4, 5])
add_para(doc, "LDraw geometry alone is not a complete connection model. For a limited MVP palette, stud and anti-stud connectivity can be generated analytically. For later expansion, the separate LDCad Shadow Library adds snapping metadata to corresponding LDraw files and is an important candidate source, subject to its share-alike license and coverage checks.", refs=[6])
add_para(doc, "BrickLink Studio is the most practical downstream compatibility target. Studio imports LDraw files, exports parts lists in several formats, can create wanted lists, and already includes a Sculpture feature that converts OBJ or STL models layer by layer using selected brick sizes, wall thickness, base thickness, colors, and optional instruction steps. That feature is a strong baseline for product quality, although it is not an embeddable service API.", refs=[7, 8])
add_para(doc, "Rebrickable is useful for catalog joins: part names, categories, external IDs, historical colors, and set inventories. BrickLink’s API adds catalog details and price-guide data. Neither should be the geometry source of truth; use them to enrich and validate a versioned local catalog derived from LDraw.", refs=[9, 10])

add_table(doc,
    ["Source", "Use in the product", "Important limitation"],
    [
        ["LDraw", "Canonical part meshes, color definitions, LDR or MPD export", "Geometry does not guarantee complete connection or market availability metadata"],
        ["LDCad Shadow Library", "Snap and connector metadata for broader palettes", "Separate coverage and CC BY SA obligations require review"],
        ["Rebrickable", "Part ID mapping, color history, inventory joins", "API is not a geometry or live marketplace feed"],
        ["BrickLink", "Price guide, buying workflow, Studio interoperability", "Authentication, marketplace variance, and service terms"],
        ["BrickLink Studio", "Reference baseline, validation, instructions, wanted lists", "Desktop product rather than a backend conversion API"],
    ], widths=[1.3, 2.75, 2.7], font_size=9)

add_heading(doc, "Image and geometry reconstruction", 2)
add_para(doc, "Two complementary families of tools are mature enough for experiments. Classical Structure from Motion and Multi View Stereo pipelines such as COLMAP recover camera poses, depth, dense point clouds, and meshes from overlapping views. They can be geometrically faithful when capture is controlled, but need texture, consistent illumination, and a mostly static object.", refs=[18])
add_para(doc, "Generative single-image systems fill hidden geometry with learned priors. SAM 3D Objects targets full shape, texture, pose, and layout from real-world images, including clutter and occlusion. TRELLIS 2 produces PBR-ready 3D assets from a single image using a sparse voxel representation. These tools are useful fallbacks and proposal generators, but their unseen surfaces are plausible estimates rather than measurements.", refs=[16, 17])
add_para(doc, "A recent open-source project with almost the same product goal, BrickBuilder AI, reports a voxel-first SAM 3D or TRELLIS pipeline followed by brick optimization and instruction output. It is useful as an engineering reference and competitive baseline, not as proof that all objects or outputs are physically reliable.", refs=[20])

add_heading(doc, "Prior work on buildable brick structures", 2)
add_para(doc, "The literature strongly supports an intermediate volumetric target followed by constrained brick placement. Constructable Brick Sculptures starts from 1 by 1 cells, greedily merges them into larger legal bricks, then uses a connection graph to repair disconnected components and weak articulation points. Legolization adds force-based stability analysis and iterative repair while accounting for shape, color, and brick count.", refs=[11, 12])
add_para(doc, "Image2Lego uses the same broad decomposition: infer a voxelized object, deterministically group voxels into larger bricks, quantize colors, and export LDraw. It also documents two recurring problems: occluded geometry from a single view and hollow shells that are difficult to build. Its authors explicitly identify multiple input images as a natural remedy for hidden geometry.", refs=[13])
add_para(doc, "Newer generative approaches are valuable but not yet a reason to remove the solver. LegoGPT trains on stable designs generated on a 20 by 20 by 20 grid using only eight common brick sizes, then adds rejection sampling and physics-aware rollback. BrickAnything conditions on point clouds and adds structure-aware tokenization, constrained decoding, and stability-guided rollback. Both show that learned proposal generation works best when paired with explicit validity mechanisms.", refs=[14, 15])

add_heading(doc, "Recommended system architecture", 1)
add_para(doc, "The architecture should preserve uncertainty until the user chooses a design. A clean mesh is not the contract between stages. The contract is a Brick Design Target containing an oriented occupancy or signed-distance field, per-view silhouettes, landmark constraints, a limited surface-color field, confidence values, and the user’s size, piece-count, price, and style preferences.")

add_table(doc,
    ["Stage", "Primary operation", "Output and gate"],
    [
        ["1 Capture", "Guide 6 to 12 views around a static object plus top, underside if safe, and one scale reference", "Coverage score; request missing angles before expensive work"],
        ["2 Understand", "Segment object, classify material and geometry risks, mark identity-defining features", "Structured reconstruction brief with uncertainties"],
        ["3 Reconstruct", "Fuse multi-view geometry; use a generative prior only to complete uncertain regions", "Normalized proxy geometry and camera agreement report"],
        ["4 Discretize", "Sample an anisotropic stud by stud by plate lattice at several candidate scales", "Occupancy, surface color, landmarks, and silhouette targets"],
        ["5 Assemble", "Cover occupied cells using a curated legal part palette and allowed orientations", "Part placements with zero collisions and legal connections"],
        ["6 Repair", "Stagger seams, connect components, support overhangs, add base and internal ribs", "Connectivity and stability checks pass"],
        ["7 Refine", "Optimize visible surfaces and landmarks with slopes, tiles, and approved templates", "Improved likeness without invalidating structure"],
        ["8 Deliver", "Generate LDR or MPD, bill of materials, preview renders, and build steps", "Studio import plus automated and human QA"],
    ], widths=[0.86, 3.25, 2.7], font_size=8.6)

add_heading(doc, "The assembly solver", 2)
add_para(doc, "Start with a deterministic optimizer. Learned systems can later propose local motifs or whole layouts, but a catalog-constrained solver is easier to debug, benchmark, and trust. A useful implementation is hierarchical rather than one enormous search over every catalog element.")
add_number(doc, 1, "Generate a target at several scales and orientations. Score each against silhouette retention, minimum wall thickness, expected piece count, and unsupported detail. Keep the best few, not just one.")
add_number(doc, 2, "Fill occupied cells with 1 by 1 primitives, then merge into larger bricks and plates. Prefer common elements and seam overlap with adjacent layers. Use randomized restarts or beam search to avoid a single greedy layout.", refs=[11, 14])
add_number(doc, 3, "Build a directed connection graph from the base upward. Reject floating components, collisions, illegal stud engagement, inaccessible insertions, and steps that require later pieces to pass through existing geometry.")
add_number(doc, 4, "Repair weak regions by splitting and remerging local bricks, rotating seams, adding concealed ties or internal columns, and simplifying unsupported target geometry. Run a lightweight graph heuristic on every candidate and a force-based check on finalists.", refs=[11, 12, 14])
add_number(doc, 5, "Quantize colors in a perceptual color space, then intersect each requested color with known part-color availability. Optimize salient surface regions first; allow hidden structural colors to come from a cheaper internal palette.")
add_number(doc, 6, "Perform a localized detail pass. Candidate substitutions—slopes, curved slopes, wedges, tiles, clips, bars, or SNOT modules—must pass the same collision, connectivity, availability, and sequence validators.")

add_para(doc, "A practical objective function is a weighted sum of silhouette error across input views, landmark displacement, occupied-volume error, surface-color error, estimated price, part count, unique part-color count, exposed studs, weak connections, and assembly complexity. Hard constraints should cover collisions, supported placement, catalog legality, allowed orientations, target size, and minimum connection rules. The user can select presets such as recognizable, low-cost, compact, or smooth.")

add_heading(doc, "Astra’s role", 1)
add_para(doc, "Astra is well suited to orchestrating this pipeline. Official OpenAI documentation describes image input, structured outputs, function calling, computer use, shell or code tools, and long multistep workflows. It does not list native 3D geometry as an output modality: the model returns text, while 3D work happens through software and tools.", refs=[1])
add_para(doc, "OpenAI’s architectural visualization example is instructive. Astra authored editable Blender scenes through Blender’s Python API, ran background renders, inspected previews, repaired geometry and shading, and exported data to Unreal. This demonstrates agentic 3D workflow control, iteration, and visual QA—not a certified photo-to-mesh or buildability engine.", refs=[2])

add_table(doc,
    ["Use Astra for", "Keep deterministic or specialized"],
    [
        ["Image inspection, object segmentation instructions, capture feedback, landmark identification, and ambiguity questions", "Segmentation masks and camera geometry produced by specialized vision or reconstruction models"],
        ["Producing a structured design brief and choosing sensible scale, style, and palette proposals", "Voxelization, part geometry, collision tests, connector compatibility, and inventory joins"],
        ["Calling reconstruction and optimization tools, tracking variants, diagnosing failures, and deciding what to retry", "The core layout search, hard constraints, structural checks, and reproducible scoring"],
        ["Rendering multiple views in Blender, comparing them with source images, and explaining tradeoffs to the user", "Final acceptance gates, Studio import tests, BOM totals, and build sequence validation"],
    ], widths=[3.35, 3.45], font_size=9)

add_para(doc, "For production, have Astra operate against narrow tools with structured schemas: analyze_capture, reconstruct_proxy, generate_target_grid, solve_layout, validate_structure, estimate_parts, render_views, and revise_region. Store every tool input, output, model version, random seed, and validator result. A model-generated design should never bypass validators because it looks convincing in a render.")

add_heading(doc, "MVP scope", 1)
add_para(doc, "A narrow MVP can prove the central algorithm without pretending to solve unrestricted custom set design. The recommended product promise is: from a guided photo set of a compact, opaque, static object, produce two or three recognizable tabletop sculpture variants that import into BrickLink Studio and use real, color-valid elements.")

add_table(doc,
    ["Include in the MVP", "Defer"],
    [
        ["Opaque static objects with a single dominant body and simple protrusions", "Transparent, reflective, articulated, very thin, furry-detail, or moving subjects"],
        ["Maximum dimension around 20 to 32 studs; roughly 100 to 800 parts", "Large sculptures, life-size models, and heavy cantilevers"],
        ["Curated bricks and plates first; a small set of slopes and tiles after the core passes", "Thousands of specialized, decorated, flexible, motorized, or Technic elements"],
        ["Upright layer-based construction on a base or small plinth", "Arbitrary-angle freeform assemblies and advanced SNOT everywhere"],
        ["Four to eight exterior colors plus hidden structural colors", "Perfect texture reproduction, stickers, custom printing, or unrestricted palettes"],
        ["LDR or MPD, BOM, preview, confidence notes, and coarse instruction steps", "LEGO-style polished instruction books and one-click global purchasing"],
    ], widths=[3.4, 3.4], font_size=9)

add_heading(doc, "Implementation phases", 2)
add_heading(doc, "Phase 0 Baselines and data", 3)
add_bullet(doc, "Create a versioned local catalog from LDraw and Rebrickable mappings. Curate 20 to 40 part families and generate analytic connectors, collision boxes, footprints, heights, legal rotations, colors, and availability scores.")
add_bullet(doc, "Build a baseline around known meshes and compare the in-house voxel-to-brick solver against BrickLink Studio Sculpture on a controlled object set.", refs=[7])
add_bullet(doc, "Select 30 to 50 benchmark objects across bottles, mugs, shoes, toys, simple animals or figurines, and household shapes. Capture ground-truth multi-view images and manually rate identity landmarks.")

add_heading(doc, "Phase 1 Mesh to buildable assembly", 3)
add_bullet(doc, "Do not start with image reconstruction. First prove that a clean mesh can become a valid LDraw model with the intended piece count, stability, palette, and visual score.")
add_bullet(doc, "Implement anisotropic voxelization, greedy or beam merging, connectivity graph analysis, seam staggering, structural repairs, BOM generation, and Studio import tests.")
add_bullet(doc, "Physically build at least ten outputs. Digital connectivity is necessary but cannot reveal every clutch, flex, tolerance, or awkward assembly issue.")

add_heading(doc, "Phase 2 Guided images to target geometry", 3)
add_bullet(doc, "Add capture guidance and multi-view reconstruction. Use SAM 3D Objects or TRELLIS 2 as a fallback or completion prior and compare against COLMAP-style reconstruction for controlled captures.", refs=[16, 17, 18])
add_bullet(doc, "Project generated brick models into every source camera and optimize silhouette and landmark agreement. Let users correct orientation, scale, a few landmarks, and uncertain hidden regions before final solving.")

add_heading(doc, "Phase 3 Visual detail and product loop", 3)
add_bullet(doc, "Add local templates and slopes only after structural metrics are stable. Use Astra to identify high-value regions and propose bounded substitutions, then validate them deterministically.")
add_bullet(doc, "Offer variants that expose the tradeoff: fewer parts, closer likeness, lower estimated price, or smoother surfaces. Collect edits and physical-build feedback as training and evaluation data.")

add_heading(doc, "Evaluation and go or no go criteria", 1)
add_para(doc, "The project should be evaluated as a physical-design system, not an image generator. Every result needs machine metrics, source-view comparisons, and a smaller physical build audit.")
add_table(doc,
    ["Dimension", "Suggested MVP metric", "Initial acceptance target"],
    [
        ["Legality", "Collision-free, valid part IDs and rotations, allowed part-color pairs", "100 percent on released designs"],
        ["Connectivity", "One base-connected component; no floating elements; minimum stud overlap rules", "100 percent on released designs"],
        ["Assembly", "Bottom-up sequence with no blocked insertion in the MVP connection model", "100 percent on released designs"],
        ["Recognition", "Blind human identification and preference versus Studio Sculpture baseline", "At least 80 percent identified; win or tie on most benchmark objects"],
        ["View fidelity", "Mean silhouette intersection over union across held-out views plus landmark error", "Set after a 30-object baseline; report by object class"],
        ["Practicality", "Piece count, unique lots, estimated price, compute time, and manual corrections", "Median under 500 parts and under ten user corrections"],
        ["Physical audit", "Independent builder completes model and handles it normally", "At least 8 of 10 without redesign"],
    ], widths=[1.0, 3.45, 2.35], font_size=8.7)

add_para(doc, "The most important early go or no go test is Phase 1. If the solver cannot outperform or materially differentiate from Studio Sculpture on buildability, controllability, or user-facing variants when given clean meshes, better image-to-3D will not rescue the product. Conversely, a strong solver can accept improving reconstruction models over time.")

add_heading(doc, "Risks and mitigations", 1)
add_table(doc,
    ["Risk", "Consequence", "Mitigation"],
    [
        ["Hidden geometry is hallucinated", "A plausible but incorrect rear or underside drives bad brick placement", "Show uncertainty, request missing views, use symmetry priors, and let users approve the proxy"],
        ["Full catalog search is intractable", "Slow solver and fragile assemblies", "Curated palette, hierarchical search, local templates, and explicit allowed orientations"],
        ["Digital model is not physically sound", "Users buy parts for an unbuildable result", "Graph and force checks, insertion checks, Studio validation, and routine physical audits"],
        ["Rare part-color combinations", "High cost or impossible procurement", "Availability-aware palette, internal hidden colors, substitutions, and price-sensitive variants"],
        ["Fidelity is lost at small scales", "Outputs look generic", "Multi-scale proposals, semantic landmarks, silhouette scoring, and honest minimum-size warnings"],
        ["License or trademark misuse", "Launch or branding risk", "Track source licenses, attribute LDraw, avoid official logos or implied endorsement, and obtain counsel before commercialization"],
    ], widths=[1.6, 2.45, 2.75], font_size=8.8)

add_para(doc, "The legal issue is not an afterthought. LDraw’s licenses can support reuse when their conditions are followed, but LEGO’s public Fair Play guidance is aimed at noncommercial fan references and warns against trademark use that implies sponsorship, use of the logo, and use of LEGO in domain names. A commercial launch should use independent branding, accurate nominative references, a clear non-endorsement statement, and specialist legal review.", refs=[5, 19])

add_heading(doc, "Immediate decisions", 1)
add_para(doc, "The following choices are sufficient to define the next project brief:")
add_number(doc, 1, "Adopt the hybrid architecture and make the catalog-constrained assembly solver the durable core product asset.")
add_number(doc, 2, "Require guided multi-view capture for the best-quality path, while allowing a clearly labeled single-image experimental mode.")
add_number(doc, 3, "Optimize the MVP for recognizable, stable tabletop sculptures rather than arbitrary expert-level LEGO models.")
add_number(doc, 4, "Use LDraw as canonical geometry and exchange, Rebrickable and BrickLink as enrichment layers, and BrickLink Studio as an external compatibility baseline.")
add_number(doc, 5, "Use Astra to plan, orchestrate, inspect, render, and converse. Keep geometry reconstruction, constraint solving, and structural acceptance in specialized tools.")
add_number(doc, 6, "Prove mesh-to-buildability before investing heavily in the upload experience or training a direct image-to-brick model.")

add_heading(doc, "Suggested first technical spike", 2)
add_para(doc, "A two-week spike should take ten clean OBJ or STL objects, voxelize each at 20, 24, and 32 studs on the longest axis, and produce LDR files using only eight to twelve common rectangular bricks and plates. It should run connectivity and seam checks, render source-aligned comparisons, import every result into Studio, and physically build two difficult examples. The output is an evidence table, not a demo reel: fidelity, parts, unique lots, solve time, validator failures, and manual repairs for every candidate.")
add_para(doc, "If that spike works, the next spike compares multi-view reconstruction, SAM 3D Objects, and TRELLIS 2 on the same captured objects after all outputs are reduced to the same Brick Design Target. This isolates whether reconstruction quality actually changes the brick result instead of rewarding a prettier intermediate mesh.", refs=[16, 17, 18])

add_heading(doc, "Conclusion", 1)
add_para(doc, "The available ecosystem is strong enough to build this product without creating a universal brick-model database from scratch. The geometry library, exchange standard, catalog metadata, desktop reference implementation, and decades of computational brick-layout research already exist. The defensible engineering work is the uncertainty-aware bridge from images to a target shape and the optimizer that turns that target into a recognizable, obtainable, stable, and sequenced assembly.")
add_para(doc, "Astra can make the system far more capable and usable by understanding the images, directing specialized tools, diagnosing failures, and iterating through rendered evidence. The product should rely on Astra’s judgment where interpretation is useful and on deterministic checks where a purchased pile of parts must fit together.")

add_heading(doc, "Notes and sources", 1)
add_para(doc, "Sources were accessed in September 2026. Product capabilities, catalog counts, pricing, APIs, and licenses can change; verify them again before implementation or launch.")
for number, (author, title_text, date, url) in SOURCES.items():
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    lead = p.add_run(f"{number}. {author}. ")
    lead.bold = True
    add_hyperlink(p, f"{title_text}.", url)
    p.add_run(f" {date}.")

core = doc.core_properties
core.title = "Image to Buildable Brick Model Architecture"
core.subject = "Research assessment and recommended MVP pipeline"
core.author = "OpenAI"
core.keywords = "brick model, image to 3D, LDraw, Astra, buildability, optimization"

doc.save(OUTPUT)
print(OUTPUT)
